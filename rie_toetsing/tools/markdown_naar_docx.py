"""Zet een Markdown-bestand om naar een Word-document (.docx).

Beperkte, doelgerichte omzetter voor de projectdocumentatie (koppen, alinea's,
opsommingen, tabellen, codeblokken, citaten en inline **vet**/`code`).

Gebruik:

    python tools/markdown_naar_docx.py HANDLEIDING.md HANDLEIDING.docx
"""

from __future__ import annotations

import re
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

AMM_BLAUW = "5467B0"
CODE_FILL = "F2F2F2"
FONT = "Aptos"
MONO = "Consolas"

_INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def _shade(par, hexcolor: str) -> None:
    """Geef een alinea een achtergrondkleur (voor codeblokken)."""
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    ppr.append(shd)


def _voeg_inline_toe(par, tekst: str, *, basis_mono: bool = False) -> None:
    """Voeg tekst toe met ondersteuning voor **vet** en `code`."""
    for deel in _INLINE.split(tekst):
        if not deel:
            continue
        run = par.add_run()
        if deel.startswith("**") and deel.endswith("**"):
            run.text = deel[2:-2]
            run.bold = True
            run.font.name = FONT
        elif deel.startswith("`") and deel.endswith("`"):
            run.text = deel[1:-1]
            run.font.name = MONO
        else:
            run.text = deel
            run.font.name = MONO if basis_mono else FONT


def _kop(doc, tekst: str, niveau: int) -> None:
    par = doc.add_heading(level=niveau)
    run = par.add_run(re.sub(r"[*`]", "", tekst))
    run.font.name = FONT
    run.font.color.rgb = RGBColor.from_string(AMM_BLAUW)
    run.font.size = Pt({1: 16, 2: 13, 3: 11}.get(niveau, 12))
    run.bold = True


def _tabel(doc, rijen: list[list[str]]) -> None:
    if not rijen:
        return
    tabel = doc.add_table(rows=0, cols=len(rijen[0]))
    tabel.style = "Table Grid"
    for i, rij in enumerate(rijen):
        cellen = tabel.add_row().cells
        for cel, waarde in zip(cellen, rij):
            par = cel.paragraphs[0]
            _voeg_inline_toe(par, waarde.strip())
            if i == 0:
                _shade(par, AMM_BLAUW)
                for run in par.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def converteer(md_pad: str, docx_pad: str) -> None:
    regels = open(md_pad, encoding="utf-8").read().splitlines()
    doc = docx.Document()
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(10)

    i = 0
    while i < len(regels):
        regel = regels[i]

        # Codeblok ```
        if regel.startswith("```"):
            i += 1
            code: list[str] = []
            while i < len(regels) and not regels[i].startswith("```"):
                code.append(regels[i])
                i += 1
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = par.add_run("\n".join(code))
            run.font.name = MONO
            run.font.size = Pt(9)
            _shade(par, CODE_FILL)
            i += 1
            continue

        # Tabel (één of meer aaneengesloten | ... | regels)
        if regel.strip().startswith("|"):
            blok: list[str] = []
            while i < len(regels) and regels[i].strip().startswith("|"):
                blok.append(regels[i])
                i += 1
            rijen = [
                [c for c in r.strip().strip("|").split("|")]
                for r in blok
                if not re.match(r"^\s*\|[\s|:-]+\|\s*$", r)
            ]
            _tabel(doc, rijen)
            continue

        gestript = regel.strip()
        if gestript.startswith("### "):
            _kop(doc, gestript[4:], 3)
        elif gestript.startswith("## "):
            _kop(doc, gestript[3:], 2)
        elif gestript.startswith("# "):
            _kop(doc, gestript[2:], 1)
        elif gestript == "---":
            pass  # horizontale lijn overslaan
        elif gestript.startswith("- "):
            par = doc.add_paragraph(style="List Bullet")
            _voeg_inline_toe(par, gestript[2:])
        elif gestript.startswith("> "):
            par = doc.add_paragraph()
            _voeg_inline_toe(par, gestript[2:])
            for run in par.runs:
                run.italic = True
        elif gestript == "":
            pass
        else:
            par = doc.add_paragraph()
            _voeg_inline_toe(par, gestript)

        i += 1

    doc.save(docx_pad)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Gebruik: python tools/markdown_naar_docx.py <in.md> <uit.docx>")
        raise SystemExit(2)
    converteer(sys.argv[1], sys.argv[2])
    print(f"Geschreven: {sys.argv[2]}")
