"""Word-export (.docx) van het Toets- en adviesrapport in de AMM-huisstijl.

Huisstijl conform het vastgestelde sjabloon (v2.2):
- primaire kleur blauw #5467B0
- lettertype Aptos (algemeen 9,5 pt, in tabellen 8 pt)
- voorpagina en samenvatting staand, tabelpagina's liggend
- RAG-codering van de oordelen (Ja groen, Deels oranje, Nee rood, n.v.t. grijs)
  met een wit vinkje in de gekleurde cel
"""

from __future__ import annotations

from datetime import date
from typing import Any

from .criteria import (
    ACTUALITEIT,
    ACTUELE_INZICHTEN,
    BETROUWBAARHEID,
    PLAN_VAN_AANPAK,
    VOLLEDIGHEID,
    Criterium,
)
from .prompt import AKD_GEGEVENS

_GEEN = "[in te vullen]"

# Huisstijlkleuren.
AMM_BLAUW = "5467B0"
RAG_FILL = {"Ja": "2E7D32", "Deels": "ED6C02", "Nee": "C62828", "n.v.t.": "9E9E9E"}

FONT = "Aptos"
FONT_TEKST_PT = 9.5
FONT_TABEL_PT = 8.0


def _imports():
    try:
        import docx  # type: ignore
        from docx.enum.section import WD_ORIENT, WD_SECTION  # type: ignore
        from docx.oxml import OxmlElement  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        from docx.shared import Pt, RGBColor  # type: ignore
    except ImportError as exc:  # pragma: no cover - afhankelijk van omgeving
        raise RuntimeError(
            "Voor de Word-export is 'python-docx' nodig: pip install python-docx"
        ) from exc
    return docx, WD_ORIENT, WD_SECTION, OxmlElement, qn, Pt, RGBColor


def _shade(cell, hexcolor: str, qn, OxmlElement) -> None:
    """Geef een tabelcel een achtergrondkleur."""
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def _set_font(run, Pt, RGBColor, *, size=FONT_TEKST_PT, bold=False, kleur=None, wit=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if wit:
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    elif kleur:
        run.font.color.rgb = RGBColor.from_string(kleur)


def _leeg_body(doc, qn) -> None:
    """Verwijder de bestaande inhoud uit een sjabloon, behoud de sectie-instellingen.

    De body-level ``w:sectPr`` (met marges, paginaformaat en verwijzingen naar
    kop-/voettekst, inclusief het logo) blijft staan, zodat de huisstijl van het
    sjabloon behouden blijft. Alleen alinea's en tabellen worden verwijderd.
    """
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def _landscape(section, WD_ORIENT) -> None:
    """Zet een sectie op liggend (A4)."""
    section.orientation = WD_ORIENT.LANDSCAPE
    breedte, hoogte = section.page_height, section.page_width
    section.page_width, section.page_height = breedte, hoogte


def _tabel(doc, criteria: tuple[Criterium, ...], index: dict[str, dict], helpers) -> None:
    """Voeg één toetstabel toe met RAG-gekleurde oordeelcellen."""
    _, _, _, OxmlElement, qn, Pt, RGBColor = helpers
    koppen = ["#", "Omschrijving criteria", "Oordeel", "Bevindingen", "Advies"]
    tabel = doc.add_table(rows=1, cols=len(koppen))
    tabel.style = "Table Grid"
    for cel, kop in zip(tabel.rows[0].cells, koppen):
        _shade(cel, AMM_BLAUW, qn, OxmlElement)
        run = cel.paragraphs[0].add_run(kop)
        _set_font(run, Pt, RGBColor, size=FONT_TABEL_PT, bold=True, wit=True)

    for crit in criteria:
        resultaat = index.get(crit.id, {})
        oordeel = resultaat.get("oordeel", "")
        rij = tabel.add_row().cells
        for cel, waarde in zip(
            rij,
            [
                crit.id,
                crit.omschrijving,
                oordeel,
                resultaat.get("bevinding") or _GEEN,
                resultaat.get("advies") or _GEEN,
            ],
        ):
            run = cel.paragraphs[0].add_run(str(waarde))
            _set_font(run, Pt, RGBColor, size=FONT_TABEL_PT)
        # RAG-codering in de oordeelkolom (index 2).
        if oordeel in RAG_FILL:
            oordeel_cel = rij[2]
            _shade(oordeel_cel, RAG_FILL[oordeel], qn, OxmlElement)
            oordeel_cel.paragraphs[0].clear()
            run = oordeel_cel.paragraphs[0].add_run(f"✓ {oordeel}")
            _set_font(run, Pt, RGBColor, size=FONT_TABEL_PT, bold=True, wit=True)


def _kop(doc, tekst: str, niveau: int, Pt, RGBColor) -> None:
    """Voeg een gekleurde kop toe in de huisstijl."""
    grootte = {1: 12, 2: 11, 3: 10}.get(niveau, 11)
    par = doc.add_heading(level=niveau)
    run = par.add_run(tekst)
    _set_font(run, Pt, RGBColor, size=grootte, bold=True, kleur=AMM_BLAUW)


def _alinea(doc, tekst: str, Pt, RGBColor, *, bold=False) -> None:
    par = doc.add_paragraph()
    run = par.add_run(tekst)
    _set_font(run, Pt, RGBColor, bold=bold)


def render_docx(
    data: dict[str, Any],
    pad: str,
    *,
    toetsingsdatum: str | None = None,
    sjabloon: str | None = None,
) -> None:
    """Schrijf het Toets- en adviesrapport als .docx-bestand naar ``pad``.

    Args:
        sjabloon: optioneel pad naar een bestaand .docx-sjabloon (bijv. het
            vastgestelde AMM-sjabloon). De inhoud daarvan wordt vervangen, maar
            de huisstijl — marges, kop-/voettekst, logo en paginanummering —
            blijft behouden.
    """
    helpers = _imports()
    docx, WD_ORIENT, WD_SECTION, OxmlElement, qn, Pt, RGBColor = helpers

    profiel = data.get("organisatieprofiel", {})
    index = {t.get("id", ""): t for t in data.get("toetsingen", [])}
    org = profiel.get("organisatienaam") or _GEEN
    datum = toetsingsdatum or profiel.get("toetsingsdatum") or date.today().isoformat()

    if sjabloon:
        doc = docx.Document(sjabloon)
        _leeg_body(doc, qn)
    else:
        doc = docx.Document()
    normaal = doc.styles["Normal"].font
    normaal.name = FONT
    normaal.size = Pt(FONT_TEKST_PT)

    # --- Voorpagina (staand) ---
    _kop(doc, "Toets- en adviesrapport Risico-Inventarisatie & -Evaluatie (RI&E)", 1, Pt, RGBColor)
    _alinea(doc, f"Organisatie: {org}", Pt, RGBColor, bold=True)
    _alinea(doc, f"Datum toetsrapport: {datum}", Pt, RGBColor)
    _alinea(doc, f"Toetser: {AKD_GEGEVENS['akd']} — {AKD_GEGEVENS['registratienummer']}", Pt, RGBColor)
    _alinea(
        doc,
        f"Adviesbureau: {AKD_GEGEVENS['organisatie']} · {AKD_GEGEVENS['email']} · "
        f"{AKD_GEGEVENS['telefoon']}",
        Pt,
        RGBColor,
    )

    # --- Samenvatting (staand) ---
    _kop(doc, "Samenvatting", 2, Pt, RGBColor)
    _alinea(
        doc,
        "Op grond van artikel 5 en 14 van de Arbowet en de toetsingscriteria in "
        "Staatscourant 2024 nr. 39674 is de RI&E getoetst op volledigheid, actualiteit, "
        "actuele inzichten en betrouwbaarheid.",
        Pt,
        RGBColor,
    )
    for label, sleutel in (
        ("Oordeel volledigheid", "oordeel_volledigheid"),
        ("Oordeel actualiteit", "oordeel_actualiteit"),
        ("Oordeel betrouwbaarheid", "oordeel_betrouwbaarheid"),
        ("Advies Plan van Aanpak", "advies_plan_van_aanpak"),
    ):
        _alinea(doc, f"{label}: {data.get(sleutel) or _GEEN}", Pt, RGBColor, bold=True)
    _alinea(doc, "Nog uit te voeren verdiepende onderzoeken:", Pt, RGBColor, bold=True)
    for onderzoek in data.get("verdiepende_onderzoeken") or [_GEEN]:
        doc.add_paragraph(str(onderzoek), style="List Bullet")

    # --- Basisgegevens (staand) ---
    _kop(doc, "Basisgegevens", 2, Pt, RGBColor)
    basis = [
        ("Organisatienaam", profiel.get("organisatienaam")),
        ("Branche / SBI-code", profiel.get("branche_sbi")),
        ("Aantal medewerkers", profiel.get("aantal_medewerkers")),
        ("Datum RI&E", profiel.get("datum_rie")),
        ("Gebruikt RI&E-instrument", profiel.get("gebruikt_instrument")),
        ("Toetsingsdatum", datum),
        ("Toetser", AKD_GEGEVENS["akd"]),
        ("Registratienummer", AKD_GEGEVENS["registratienummer"]),
    ]
    basistabel = doc.add_table(rows=0, cols=2)
    basistabel.style = "Table Grid"
    for label, waarde in basis:
        rij = basistabel.add_row().cells
        run = rij[0].paragraphs[0].add_run(label)
        _set_font(run, Pt, RGBColor, bold=True)
        run = rij[1].paragraphs[0].add_run(str(waarde or _GEEN))
        _set_font(run, Pt, RGBColor)

    # --- Hoofdstuk 1 + 2 (liggend) ---
    sectie = doc.add_section(WD_SECTION.NEW_PAGE)
    _landscape(sectie, WD_ORIENT)

    _kop(doc, "1. Toetsen van de RI&E", 2, Pt, RGBColor)
    for titel, criteria in (
        ("Tabel 1.1 — Toetsen op volledigheid", VOLLEDIGHEID),
        ("Tabel 1.2 — Toetsen op actualiteit", ACTUALITEIT),
        ("Tabel 1.3 — Toetsen op actuele inzichten", ACTUELE_INZICHTEN),
        ("Tabel 1.4 — Toetsen op betrouwbaarheid", BETROUWBAARHEID),
    ):
        _kop(doc, titel, 3, Pt, RGBColor)
        _tabel(doc, criteria, index, helpers)

    _kop(doc, "Samenvatting per scope", 3, Pt, RGBColor)
    scopetabel = doc.add_table(rows=0, cols=2)
    scopetabel.style = "Table Grid"
    for scope, sleutel in (
        ("AH — Arbeidshygiëne", "scope_ah"),
        ("HVK — Hogere Veiligheidskunde", "scope_hvk"),
        ("A&O — Arbeids- en Organisatiedeskundige", "scope_ao"),
    ):
        rij = scopetabel.add_row().cells
        run = rij[0].paragraphs[0].add_run(scope)
        _set_font(run, Pt, RGBColor, size=FONT_TABEL_PT, bold=True)
        run = rij[1].paragraphs[0].add_run(str(data.get(sleutel) or _GEEN))
        _set_font(run, Pt, RGBColor, size=FONT_TABEL_PT)

    _kop(doc, "2. Toetsen van het Plan van Aanpak", 2, Pt, RGBColor)
    _tabel(doc, PLAN_VAN_AANPAK, index, helpers)

    # --- Afronding (terug naar staand) ---
    sectie = doc.add_section(WD_SECTION.NEW_PAGE)
    sectie.orientation = WD_ORIENT.PORTRAIT
    breedte, hoogte = sectie.page_height, sectie.page_width
    sectie.page_width, sectie.page_height = breedte, hoogte

    _kop(doc, "3. Afronding van de toetsing", 2, Pt, RGBColor)
    _alinea(
        doc,
        "De rapportage dient ter beschikking te worden gesteld aan de ondernemingsraad "
        "(OR) of personeelsvertegenwoordiging (PVT). De werkgever is conform artikel 14 "
        "lid 3 van de Arbowet verplicht een afschrift te zenden aan de belanghebbende "
        "werknemers.",
        Pt,
        RGBColor,
    )
    _alinea(doc, "Toetsing uitgevoerd door:", Pt, RGBColor, bold=True)
    _alinea(doc, AKD_GEGEVENS["akd"], Pt, RGBColor)
    _alinea(doc, AKD_GEGEVENS["registratienummer"], Pt, RGBColor)
    _alinea(doc, f"{AKD_GEGEVENS['organisatie']} — datum: {datum}", Pt, RGBColor)
    if data.get("samenvatting"):
        _alinea(doc, data["samenvatting"], Pt, RGBColor)
    _alinea(
        doc,
        "Dit rapport is opgesteld met ondersteuning van een digitale toetsingsassistent "
        "op basis van het certificatieschema 2024 (Stcrt. 39674), Arbowet, Arbobesluit, "
        "Arboregeling, de Steunpunt RI&E-criteria en relevante NEN-normen. De inhoudelijke "
        "eindverantwoordelijkheid berust bij de ondertekenende gecertificeerde "
        "arbokerndeskundige.",
        Pt,
        RGBColor,
    )

    doc.save(pad)
