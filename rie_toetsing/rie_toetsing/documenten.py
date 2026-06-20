"""Inlezen van de te toetsen documenten (RI&E, PvA, BHV-plan, arbobeleid).

PDF's worden als document-bron rechtstreeks aan de Claude API meegegeven; tekst-
en Word-bestanden worden naar platte tekst omgezet.
"""

from __future__ import annotations

import base64
from dataclasses import dataclass
from pathlib import Path

# Bestandstypen die we ondersteunen.
PDF_SUFFIXES = {".pdf"}
TEKST_SUFFIXES = {".txt", ".md"}
WORD_SUFFIXES = {".docx"}
ONDERSTEUND = PDF_SUFFIXES | TEKST_SUFFIXES | WORD_SUFFIXES


@dataclass(frozen=True)
class Document:
    """Een ingelezen brondocument met een label (bijv. 'RI&E' of 'Plan van Aanpak')."""

    label: str
    bestandsnaam: str
    soort: str  # "pdf" of "tekst"
    pdf_base64: str | None = None
    tekst: str | None = None


def _lees_word(pad: Path) -> str:
    """Haal de platte tekst uit een .docx-bestand."""
    try:
        from docx import Document as DocxDocument  # type: ignore
    except ImportError as exc:  # pragma: no cover - afhankelijk van omgeving
        raise RuntimeError(
            "Voor .docx-bestanden is 'python-docx' nodig: pip install python-docx"
        ) from exc

    document = DocxDocument(str(pad))
    return "\n".join(p.text for p in document.paragraphs if p.text)


def laad_document(pad: str | Path, label: str) -> Document:
    """Lees één brondocument in en geef een :class:`Document` terug."""
    pad = Path(pad)
    if not pad.is_file():
        raise FileNotFoundError(f"Bestand niet gevonden: {pad}")

    suffix = pad.suffix.lower()
    if suffix not in ONDERSTEUND:
        raise ValueError(
            f"Bestandstype '{suffix}' wordt niet ondersteund "
            f"({', '.join(sorted(ONDERSTEUND))}). Bestand: {pad.name}"
        )

    if suffix in PDF_SUFFIXES:
        data = base64.standard_b64encode(pad.read_bytes()).decode("ascii")
        return Document(label=label, bestandsnaam=pad.name, soort="pdf", pdf_base64=data)

    if suffix in WORD_SUFFIXES:
        tekst = _lees_word(pad)
    else:
        tekst = pad.read_text(encoding="utf-8", errors="replace")

    return Document(label=label, bestandsnaam=pad.name, soort="tekst", tekst=tekst)


def naar_contentblokken(documenten: list[Document]) -> list[dict]:
    """Zet documenten om naar content-blokken voor een Claude-bericht."""
    blokken: list[dict] = []
    for doc in documenten:
        if doc.soort == "pdf":
            blokken.append(
                {
                    "type": "document",
                    "title": f"{doc.label} — {doc.bestandsnaam}",
                    "source": {
                        "type": "base64",
                        "media_type": "application/pdf",
                        "data": doc.pdf_base64,
                    },
                }
            )
        else:
            blokken.append(
                {
                    "type": "text",
                    "text": (
                        f"=== {doc.label} — {doc.bestandsnaam} ===\n{doc.tekst}"
                    ),
                }
            )
    return blokken
