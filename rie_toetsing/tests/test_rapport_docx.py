"""Tests voor de Word-export (vereist python-docx)."""

import pytest

from rie_toetsing import criteria

docx = pytest.importorskip("docx")
from rie_toetsing.rapport_docx import render_docx  # noqa: E402


def _data():
    return {
        "organisatieprofiel": {"organisatienaam": "Voorbeeld B.V."},
        "toetsingen": [
            {"id": c.id, "oordeel": "Ja", "bevinding": f"Bevinding {c.id}",
             "advies": f"Advies {c.id}"}
            for c in criteria.ALLE_CRITERIA
        ],
        "oordeel_volledigheid": "Voldoende.",
        "oordeel_actualiteit": "Actueel.",
        "oordeel_betrouwbaarheid": "Betrouwbaar.",
        "advies_plan_van_aanpak": "SMART.",
        "verdiepende_onderzoeken": ["PSA-onderzoek"],
        "scope_ah": "AH.",
        "scope_hvk": "HVK.",
        "scope_ao": "A&O.",
        "samenvatting": "Eindoordeel voldoende.",
    }


def test_render_docx_maakt_leesbaar_bestand(tmp_path):
    pad = tmp_path / "rapport.docx"
    render_docx(_data(), str(pad))
    assert pad.is_file() and pad.stat().st_size > 0

    document = docx.Document(str(pad))
    tekst = "\n".join(p.text for p in document.paragraphs)
    assert "Toets- en adviesrapport" in tekst
    assert "Voorbeeld B.V." in tekst

    # Alle toetsregels moeten in de tabellen voorkomen.
    celteksten = {
        cel.text
        for tabel in document.tables
        for rij in tabel.rows
        for cel in rij.cells
    }
    plat = " ".join(celteksten)
    for crit in criteria.ALLE_CRITERIA:
        assert crit.id in plat


def test_render_docx_leeg_sjabloon(tmp_path):
    pad = tmp_path / "leeg.docx"
    render_docx({}, str(pad))
    assert pad.is_file()


def _maak_sjabloon(pad, Cm):
    """Bouw een testsjabloon met koptekst, voettekst, marge en oude inhoud."""
    sjabloon = docx.Document()
    sectie = sjabloon.sections[0]
    sectie.left_margin = Cm(3)
    sectie.header.paragraphs[0].text = "AMM CONSULTANCY B.V."
    sectie.footer.paragraphs[0].text = "Vertrouwelijk — pagina"
    sjabloon.add_paragraph("OUDE PLACEHOLDER-INHOUD die weg moet")
    sjabloon.add_table(rows=1, cols=1)
    sjabloon.save(str(pad))


def test_render_docx_met_sjabloon_behoudt_huisstijl(tmp_path):
    from docx.shared import Cm

    sjabloon_pad = tmp_path / "sjabloon.docx"
    _maak_sjabloon(sjabloon_pad, Cm)

    uit = tmp_path / "rapport.docx"
    render_docx(_data(), str(uit), sjabloon=str(sjabloon_pad))

    document = docx.Document(str(uit))
    eerste = document.sections[0]

    # Huisstijl uit het sjabloon blijft behouden.
    assert eerste.header.paragraphs[0].text == "AMM CONSULTANCY B.V."
    assert "Vertrouwelijk" in eerste.footer.paragraphs[0].text
    assert round(eerste.left_margin.cm) == 3

    # Oude placeholder-inhoud is verdwenen, nieuwe inhoud staat erin.
    body = "\n".join(p.text for p in document.paragraphs)
    assert "OUDE PLACEHOLDER-INHOUD" not in body
    assert "Toets- en adviesrapport" in body
    assert "Voorbeeld B.V." in body
